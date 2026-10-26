"""Respaldo semanal de médicos + historial completo de ventas.

- Genera un Excel (sin datos de pacientes: se excluyen rut_paciente y nombre_paciente).
- Lo guarda en la tabla data_backups (historial descargable desde Admin → Respaldos).
- Lo envía por correo como adjunto (SMTP_* + BACKUP_EMAIL, o REPORT_EMAIL si no existe).
- Corre cada lunes 6:00 hora Chile y a pedido (POST /api/backups).
"""
import io
import json
import os
import smtplib
from datetime import datetime
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

import pytz
from openpyxl import Workbook
from sqlalchemy import func, extract
from sqlalchemy.orm import Session

from .models import BusinessLine, DataBackup, Doctor, MedicalRep, Sale, Visit

CHILE_TZ = pytz.timezone("America/Santiago")
KEEP_LAST = 26                 # ~6 meses de respaldos semanales en la app
MAX_ATTACH_BYTES = 20 * 1024 * 1024


def _fmt(d):
    return d.strftime("%Y-%m-%d %H:%M") if isinstance(d, datetime) else d


def build_backup_xlsx(db: Session) -> tuple:
    lines = {b.id: b.name for b in db.query(BusinessLine).all()}
    reps = db.query(MedicalRep).order_by(MedicalRep.id).all()
    rep_names = {r.id: r.name for r in reps}
    doctors = db.query(Doctor).order_by(Doctor.id).all()
    doc_names = {d.id: d.name for d in doctors}

    wb = Workbook(write_only=True)
    counts = {}

    # Ventas agregadas por médico (para la hoja Médicos)
    agg = dict()
    for did, n, tot, first, last in db.query(
        Sale.doctor_id, func.count(Sale.id), func.sum(Sale.amount), func.min(Sale.sale_date), func.max(Sale.sale_date)
    ).filter(Sale.doctor_id.isnot(None)).group_by(Sale.doctor_id).all():
        agg[did] = (n, float(tot or 0), first, last)

    ws = wb.create_sheet("Medicos")
    ws.append(["ID", "Nombre", "RUT", "Especialidad", "Centro médico", "Dirección", "Ciudad", "Comuna",
               "Teléfono", "Correo", "Visitador ID", "Visitador", "Línea de negocio", "Frecuencia visita (días)",
               "Productos que prescribe", "Notas", "Activo", "Origen", "Fecha alta",
               "Total ventas (uds)", "Monto total", "Primera venta", "Última venta"])
    for d in doctors:
        n, tot, first, last = agg.get(d.id, (0, 0.0, None, None))
        ws.append([d.id, d.name, d.rut, d.specialty, d.medical_center, d.address, d.city, d.commune,
                   d.phone, d.email, d.rep_id, rep_names.get(d.rep_id), lines.get(d.business_line_id),
                   d.visit_frequency, d.prescribes_products, d.notes, "Sí" if d.is_active else "No",
                   d.created_from, _fmt(d.created_at), n, tot, _fmt(first), _fmt(last)])
    counts["medicos"] = len(doctors)

    ws = wb.create_sheet("Ventas")
    ws.append(["ID", "Fecha venta", "Médico ID", "Médico", "RUT médico", "Nombre médico (archivo)",
               "Producto", "Categoría", "Monto", "ID externo", "Carga ID", "Fecha registro"])
    n_sales = 0
    for s in db.query(Sale).order_by(Sale.sale_date, Sale.id).yield_per(2000):
        ws.append([s.id, _fmt(s.sale_date), s.doctor_id, doc_names.get(s.doctor_id), s.rut_doctor,
                   s.doctor_name_raw, s.product, s.categoria, s.amount, s.external_id, s.upload_id,
                   _fmt(s.created_at)])
        n_sales += 1
    counts["ventas"] = n_sales

    ws = wb.create_sheet("Ventas_por_medico_mes")
    ws.append(["Médico ID", "Médico", "Año", "Mes", "Unidades", "Monto"])
    rows = db.query(Sale.doctor_id, extract("year", Sale.sale_date), extract("month", Sale.sale_date),
                    func.count(Sale.id), func.sum(Sale.amount)) \
        .filter(Sale.doctor_id.isnot(None), Sale.sale_date.isnot(None)) \
        .group_by(Sale.doctor_id, extract("year", Sale.sale_date), extract("month", Sale.sale_date)) \
        .order_by(Sale.doctor_id, extract("year", Sale.sale_date), extract("month", Sale.sale_date)).all()
    for did, y, m, n, tot in rows:
        ws.append([did, doc_names.get(did), int(y), int(m), n, float(tot or 0)])

    ws = wb.create_sheet("Visitas")
    ws.append(["ID", "Fecha programada", "Fecha real", "Estado", "Médico ID", "Médico",
               "Visitador ID", "Visitador", "Notas", "Origen"])
    n_vis = 0
    for v in db.query(Visit).order_by(Visit.scheduled_date, Visit.id).yield_per(2000):
        ws.append([v.id, _fmt(v.scheduled_date), _fmt(v.actual_date), v.status, v.doctor_id,
                   doc_names.get(v.doctor_id), v.rep_id, rep_names.get(v.rep_id), v.notes,
                   getattr(v, "origin", None)])
        n_vis += 1
    counts["visitas"] = n_vis

    ws = wb.create_sheet("Visitadores")
    ws.append(["ID", "Nombre", "Correo", "Teléfono", "Territorio", "Zona", "Activo", "Fecha alta"])
    for r in reps:
        ws.append([r.id, r.name, r.email, r.phone, r.territory, r.zone, "Sí" if r.is_active else "No",
                   _fmt(r.created_at)])
    counts["visitadores"] = len(reps)

    now = datetime.now(CHILE_TZ)
    ws = wb.create_sheet("Resumen")
    ws.append(["Respaldo Visitadores Médicos — Narma"])
    ws.append(["Generado", now.strftime("%Y-%m-%d %H:%M"), "hora Chile"])
    ws.append(["Médicos", counts["medicos"]])
    ws.append(["Ventas", counts["ventas"]])
    ws.append(["Visitas", counts["visitas"]])
    ws.append(["Visitadores", counts["visitadores"]])
    ws.append(["Nota", "No incluye RUT ni nombre de pacientes."])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), counts


def _send_email(filename: str, data: bytes, counts: dict) -> tuple:
    host = os.getenv("SMTP_HOST", "smtp.gmail.com")
    port = int(os.getenv("SMTP_PORT", "587"))
    user = (os.getenv("SMTP_USER") or "").strip()
    pwd = (os.getenv("SMTP_PASSWORD") or "").strip()
    to = (os.getenv("BACKUP_EMAIL") or os.getenv("REPORT_EMAIL") or "").strip()
    if not (user and pwd and to):
        return False, "Correo no configurado (faltan SMTP_USER, SMTP_PASSWORD o BACKUP_EMAIL/REPORT_EMAIL)"
    msg = MIMEMultipart()
    msg["From"] = user
    msg["To"] = to
    msg["Subject"] = f"Respaldo semanal Visitadores Médicos — {datetime.now(CHILE_TZ).strftime('%d/%m/%Y')}"
    body = (f"Respaldo automático de la plataforma de visitadores.\n\n"
            f"Médicos: {counts['medicos']}\nVentas: {counts['ventas']}\nVisitas: {counts['visitas']}\n"
            f"Visitadores: {counts['visitadores']}\n\nNo incluye datos de pacientes.\n"
            f"También queda disponible en Admin → Respaldos.")
    if len(data) > MAX_ATTACH_BYTES:
        body += "\n\nEl archivo supera 20 MB, por eso no va adjunto: descárgalo desde Admin → Respaldos."
    else:
        part = MIMEApplication(data, Name=filename)
        part["Content-Disposition"] = f'attachment; filename="{filename}"'
        msg.attach(part)
    msg.attach(MIMEText(body, "plain", "utf-8"))
    try:
        with smtplib.SMTP(host, port, timeout=60) as s:
            s.ehlo()
            s.starttls()
            s.login(user, pwd)
            s.sendmail(user, [a.strip() for a in to.split(",") if a.strip()], msg.as_string())
        return True, None
    except Exception as e:
        return False, f"Error SMTP: {e}"[:500]


def create_backup(db: Session, trigger: str = "manual", send_email: bool = True) -> DataBackup:
    data, counts = build_backup_xlsx(db)
    stamp = datetime.now(CHILE_TZ).strftime("%Y-%m-%d_%H%M")
    filename = f"Respaldo_Visitadores_{stamp}.xlsx"
    emailed, err = (_send_email(filename, data, counts) if send_email else (False, None))
    b = DataBackup(filename=filename, size_bytes=len(data), data=data, trigger=trigger,
                   summary=json.dumps(counts), emailed=emailed, email_error=err)
    db.add(b)
    db.commit()
    db.refresh(b)
    # Retención: conservar solo los últimos KEEP_LAST
    old = db.query(DataBackup.id).order_by(DataBackup.created_at.desc(), DataBackup.id.desc()).offset(KEEP_LAST).all()
    if old:
        db.query(DataBackup).filter(DataBackup.id.in_([o[0] for o in old])).delete(synchronize_session=False)
        db.commit()
    return b


def run_weekly_backup_job() -> None:
    from .database import SessionLocal
    db = SessionLocal()
    try:
        create_backup(db, trigger="semanal")
    except Exception:
        db.rollback()
    finally:
        db.close()
