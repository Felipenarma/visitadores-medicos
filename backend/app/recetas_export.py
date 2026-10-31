"""Excel "Recetas por médico según visitador": mismo formato que el análisis del 25-10-2026,
generado en vivo desde la base de datos, con venta neta (sin IVA) y período opcional.

Hojas: Resumen_grupo, Resumen_visitador, Detalle_medicos, Notas.
"""
import io
from calendar import monthrange
from datetime import datetime
from statistics import median
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy import func
from sqlalchemy.orm import Session

from .constants import IVA_RATE
from .models import BusinessLine, Doctor, MedicalRep, Sale

# Visitadores cuyos médicos se consideran "Cuentas grandes"
CUENTAS_GRANDES_REPS = {"felipe"}
GROUPS = ["Cuentas grandes", "Visitador en terreno", "Sin visitador"]
MONTHS = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
          "septiembre", "octubre", "noviembre", "diciembre"]

HDR_FONT = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="0F1E2D")
CLP = '"$"#,##0'


def _group(rep_name: Optional[str]) -> str:
    if not rep_name:
        return "Sin visitador"
    return "Cuentas grandes" if rep_name.strip().lower() in CUENTAS_GRANDES_REPS else "Visitador en terreno"


def _period(desde: Optional[str], hasta: Optional[str]):
    """'YYYY-MM' → (inicio, fin, etiqueta). Sin parámetros = histórico completo."""
    start = end = None
    if desde:
        y, m = map(int, desde.split("-"))
        start = datetime(y, m, 1)
    if hasta:
        y, m = map(int, hasta.split("-"))
        end = datetime(y, m, monthrange(y, m)[1], 23, 59, 59)
    if start and end and start > end:
        start, end = datetime(end.year, end.month, 1), datetime(start.year, start.month, monthrange(start.year, start.month)[1], 23, 59, 59)
    if not start and not end:
        label = "Histórico completo: todas las recetas cargadas en la app desde el inicio."
    else:
        a = f"{MONTHS[start.month - 1]} {start.year}" if start else "el inicio"
        b = f"{MONTHS[end.month - 1]} {end.year}" if end else "hoy"
        label = f"Desde {a} hasta {b} (inclusive)."
    return start, end, label


def _header(ws, cols, widths):
    ws.append(cols)
    for i, c in enumerate(ws[ws.max_row], start=0):
        c.font, c.fill = HDR_FONT, HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = widths[i] if i < len(widths) else 16


def _stats(rows):
    """rows: lista de (recetas, neto, bruto) de médicos con recetas > 0."""
    n = len(rows)
    rec = sum(r[0] for r in rows)
    net = sum(r[1] for r in rows)
    gross = sum(r[2] for r in rows)
    return {
        "n": n, "rec": rec, "avg": rec / n if n else 0, "med": median([r[0] for r in rows]) if n else 0,
        "net": net, "gross": gross, "avg_net": net / n if n else 0,
    }


def build_recetas_xlsx(db: Session, desde: Optional[str] = None, hasta: Optional[str] = None) -> tuple:
    start, end, period_label = _period(desde, hasta)

    reps = {r.id: r.name for r in db.query(MedicalRep).all()}
    lines = {b.id: b.name for b in db.query(BusinessLine).all()}
    doctors = db.query(Doctor).filter(Doctor.is_active == True).all()  # noqa: E712

    q = db.query(Sale.doctor_id, func.count(Sale.id), func.coalesce(func.sum(Sale.amount), 0)) \
        .filter(Sale.doctor_id.isnot(None))
    if start:
        q = q.filter(Sale.sale_date >= start)
    if end:
        q = q.filter(Sale.sale_date <= end)
    agg = {}
    for did, n, tot in q.group_by(Sale.doctor_id).all():
        try:
            tot = float(tot or 0)
            tot = 0.0 if tot != tot else tot
        except Exception:
            tot = 0.0
        agg[did] = (int(n), tot)

    detail = []
    for d in doctors:
        n, gross = agg.get(d.id, (0, 0.0))
        rep_name = reps.get(d.rep_id) if d.rep_id else None
        detail.append({
            "id": d.id, "name": d.name, "rut": d.rut, "spec": d.specialty,
            "rep": rep_name or "Sin asignar", "group": _group(rep_name),
            "line": lines.get(d.business_line_id), "rec": n,
            "net": round(gross / (1 + IVA_RATE)), "gross": round(gross),
        })
    detail.sort(key=lambda x: (-x["rec"], -x["gross"], x["name"] or ""))
    prescribers = [x for x in detail if x["rec"] > 0]

    wb = Workbook()

    # ── Resumen_grupo ──
    ws = wb.active
    ws.title = "Resumen_grupo"
    scope = "en el período" if (start or end) else "al menos una vez"
    ws.append([f"Recetas por médico según tipo de asignación — solo médicos activos que han recetado {scope}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Período: {period_label}"])
    ws.append([])
    _header(ws, ["Grupo", "Médicos que recetan", "Recetas totales", "Promedio recetas por médico",
                 "Mediana recetas por médico", "Venta neta total (CLP, sin IVA)", "Venta neta promedio por médico (CLP)",
                 "% de las recetas", "Venta total con IVA (CLP)"], [26, 14, 14, 16, 16, 20, 20, 12, 20])
    by_group = {g: _stats([(x["rec"], x["net"], x["gross"]) for x in prescribers if x["group"] == g]) for g in GROUPS}
    total = _stats([(x["rec"], x["net"], x["gross"]) for x in prescribers])
    for g in GROUPS + ["Total"]:
        s = total if g == "Total" else by_group[g]
        ws.append([g, s["n"], s["rec"], round(s["avg"], 1), s["med"], s["net"], round(s["avg_net"]),
                   (s["rec"] / total["rec"]) if total["rec"] else 0, s["gross"]])
        r = ws.max_row
        for col in "FGI":
            ws[f"{col}{r}"].number_format = CLP
        ws[f"H{r}"].number_format = "0.0%"
        if g == "Total":
            for c in ws[r]:
                c.font = Font(bold=True)
    ws.append([])
    ws.append(["Comparación clave"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    t, s = by_group["Visitador en terreno"], by_group["Sin visitador"]
    ws.append(["Veces más recetas: visitador en terreno vs sin visitador (promedio)", round(t["avg"] / s["avg"], 1) if s["avg"] else None])
    ws.append(["Veces más recetas: visitador en terreno vs sin visitador (mediana)", round(t["med"] / s["med"], 1) if s["med"] else None])
    ws.append(["Veces más venta: visitador en terreno vs sin visitador (promedio por médico)", round(t["avg_net"] / s["avg_net"], 1) if s["avg_net"] else None])

    # ── Resumen_visitador ──
    ws = wb.create_sheet("Resumen_visitador")
    _header(ws, ["Visitador", "Grupo", "Médicos asignados (activos)", "Médicos que recetan", "Recetas totales",
                 "Promedio recetas por médico", "Mediana recetas por médico", "Venta neta total (CLP, sin IVA)",
                 "Venta neta promedio por médico (CLP)", "Venta total con IVA (CLP)"], [24, 20, 14, 14, 14, 16, 16, 20, 20, 20])
    rep_rows = []
    for rep_name in sorted({x["rep"] for x in detail}):
        assigned = [x for x in detail if x["rep"] == rep_name]
        s = _stats([(x["rec"], x["net"], x["gross"]) for x in assigned if x["rec"] > 0])
        rep_rows.append((rep_name, assigned[0]["group"], len(assigned), s))
    rep_rows.sort(key=lambda r: (r[0] == "Sin asignar", -r[3]["net"]))
    for rep_name, grp, n_assigned, s in rep_rows:
        ws.append([rep_name, grp, n_assigned, s["n"], s["rec"], round(s["avg"], 1), s["med"], s["net"], round(s["avg_net"]), s["gross"]])
        for col in "HIJ":
            ws[f"{col}{ws.max_row}"].number_format = CLP

    # ── Detalle_medicos ──
    ws = wb.create_sheet("Detalle_medicos")
    _header(ws, ["ID médico", "Médico", "RUT", "Especialidad", "Visitador", "Grupo", "Línea de negocio", "Recetas",
                 "Venta neta (CLP, sin IVA)", "Venta con IVA (CLP)", "Prescriptor"], [10, 34, 14, 22, 22, 20, 20, 10, 18, 18, 11])
    for x in detail:
        ws.append([x["id"], x["name"], x["rut"], x["spec"], x["rep"], x["group"], x["line"], x["rec"],
                   x["net"], x["gross"], "Sí" if x["rec"] > 0 else "No"])
        r = ws.max_row
        ws[f"I{r}"].number_format = CLP
        ws[f"J{r}"].number_format = CLP
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:K{ws.max_row}"

    # ── Notas ──
    ws = wb.create_sheet("Notas")
    _header(ws, ["Campo", "Detalle"], [22, 110])
    terreno = sorted({x["rep"] for x in detail if x["group"] == "Visitador en terreno"})
    grandes = sorted({x["rep"] for x in detail if x["group"] == "Cuentas grandes"})
    for row in [
        ("Fuente", "Base de datos de la App Visitadores Médicos, generado en el momento de la descarga."),
        ("Fecha de corte", datetime.now().strftime("%d-%m-%Y %H:%M")),
        ("Período", period_label),
        ("Universo", f"{len(detail)} médicos activos; los resúmenes consideran solo los {len(prescribers)} que recetaron {scope}."),
        ("Receta", "Cada venta registrada en la app asociada a la ficha del médico."),
        ("Venta neta", f"Monto sin IVA = monto cargado / {1 + IVA_RATE:.2f}. Los montos cargados son precio al público con IVA."),
        ("Cuentas grandes", f"Médicos asignados a {', '.join(grandes) or 'Felipe'}, que lleva las cuentas grandes."),
        ("Visitador en terreno", f"Médicos asignados a {', '.join(terreno) or '—'}."),
        ("Sin visitador", "Médicos activos con ficha en el CRM y sin visitador asignado."),
        ("Visitador considerado", "El asignado hoy, no el que tenía el médico al momento de cada receta."),
        ("Exclusiones", "No incluye ventas sin ficha de médico en el CRM ni médicos inactivos."),
        ("Interpretación", "Muestra relación, no causalidad: los visitadores pueden estar asignados a médicos que ya recetaban más."),
    ]:
        ws.append(list(row))
        ws.cell(ws.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    suffix = "historico" if not (start or end) else \
        f"{start.strftime('%Y-%m') if start else 'inicio'}_a_{end.strftime('%Y-%m') if end else 'hoy'}"
    return buf.getvalue(), f"Recetas_por_medico_segun_visitador_{suffix}.xlsx"
